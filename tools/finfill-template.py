#!/usr/bin/env python3
"""Converts a valuation workbook (.xlsx exported from Google Sheets) into a FinFill template JSON.

    python3 tools/finfill-template.py "NVDA Valuation - 5y.xlsx" \
        --universal \
        --json Packages/FinFill/Sources/FinFillCore/Resources/templates/valuation-universal.json \
        --xlsx valuation-universal.xlsx

--universal turns the NVDA workbook into a template for any company:
  * statements (IncomeStatement / BalanceSheets / CashFlowStatement): 9 year columns G..O,
    typed-in numbers removed, formulas kept, header G2 = {{FY_FIRST}}, H2.. = previous + 1,
  * DCF / WACC look the years up with INDEX/MATCH (the model rolls forward when a year column
    is added), NVDA-specific assumptions replaced by history-based formulas / GOOGLEFINANCE,
  * placeholders {{TICKER}}, {{FY_FIRST}}, {{TODAY}}, {{FYEND}} filled in by the app.

The JSON keeps values, formulas and formatting (fonts, fills, number formats, borders,
alignment, merges, column widths, hidden gridlines) — the app recreates the workbook with the
Google Sheets API (spreadsheets.create + updateCells).
"""
import argparse, json, re
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

STATEMENTS = ["IncomeStatement", "BalanceSheets", "CashFlowStatement"]
YEAR_FIRST, YEAR_COUNT, HEADER_ROW = 7, 9, 2            # G..O (1-based), header in row 2
MAX_ROWS, MAX_COLS = 70, 26
DEFAULT_WIDTH_PX = 100                                   # Google Sheets default column width
EXCEL_DEFAULT_WIDTH = 12.63                              # = 100 px in the export

HAS_REF = re.compile(r'(?<![A-Za-z_])\$?[A-Za-z]{1,3}\$?[0-9]+')


def refs_cells(formula: str) -> bool:
    return bool(HAS_REF.search(re.sub(r'"[^"]*"', '', formula)))


def has_typed_numbers(formula: str) -> bool:
    """=(1112+6081+4500)-M12 or =6682-M21: numbers typed by hand for one year."""
    rest = HAS_REF.sub('', re.sub(r'"[^"]*"', '', formula))
    return bool(re.search(r'(?<![A-Za-z_.])(\d{2,}|\d+\.\d+)', rest))


def unwrap_google(v):
    """Google exports its own functions as IFERROR(__xludf.DUMMYFUNCTION("…"), cached)."""
    if isinstance(v, str) and '__xludf.DUMMYFUNCTION' in v:
        m = re.search(r'__xludf\.DUMMYFUNCTION\("((?:[^"]|"")*)"\)', v)
        if m:
            return '=' + m.group(1).replace('""', '"')
    return v


def idx(sheet, row):
    """INDEX/MATCH of a statement row for the year in the same DCF column (row 21)."""
    return f'INDEX({sheet}!$G${row}:$Z${row},MATCH({{col}}$21,{sheet}!$G$2:$Z$2,0))'


def universalize(wb):
    del_sheets = [n for n in wb.sheetnames if n not in STATEMENTS + ["DCF", "WACC"]]
    for n in del_sheets:
        del wb[n]

    for name in STATEMENTS:
        ws = wb[name]
        for c in range(YEAR_FIRST, YEAR_FIRST + YEAR_COUNT):
            for r in range(HEADER_ROW + 1, MAX_ROWS + 1):
                cell = ws.cell(r, c)
                v = cell.value
                if v is None:
                    continue
                if isinstance(v, str) and v.startswith('=') and refs_cells(v) and not has_typed_numbers(v):
                    continue            # totals, ratios, NWC … stay
                cell.value = None       # typed-in number or constant formula (=6283+1793)
            letter = get_column_letter(c)
            prev = get_column_letter(c - 1)
            ws.cell(HEADER_ROW, c).value = '{{FY_FIRST}}' if c == YEAR_FIRST else f'={prev}{HEADER_ROW}+1'
    wb["IncomeStatement"]["B3"] = "Revenue"
    wb["IncomeStatement"]["B4"] = "Other revenue"

    d = wb["DCF"]
    d["B1"] = '=C4&" DCF"'
    d["C4"] = "{{TICKER}}"
    d["C5"] = "={{TODAY}}"
    d["C6"] = "={{FYEND}}"
    d["H5"] = "=GOOGLEFINANCE(C4)"
    for a in ("K4", "K5", "K6", "J4", "J5", "J6"):
        d[a] = None
    # Base case everywhere (1 conservative, 2 base, 3 optimistic).
    for a in ("C12", "C13", "C14", "C15"):
        d[a] = 2
    # Years: the last five actuals, found from the statements.
    d["D21"] = "=MAX(IncomeStatement!$G$2:$Z$2)-4"
    hist = {22: ("IncomeStatement", 5, ""), 28: ("IncomeStatement", 11, ""), 34: ("IncomeStatement", 14, ""),
            39: ("CashFlowStatement", 4, ""), 42: ("CashFlowStatement", 15, "-"), 45: ("BalanceSheets", 38, "")}
    for row, (sheet, srow, sign) in hist.items():
        for col in "DEFGH":
            d[f"{col}{row}"] = "=" + sign + idx(sheet, srow).replace("{col}", col)
    last = lambda sheet, r: idx(sheet, r).replace("{col}", "$H")
    d["M54"] = f"={last('BalanceSheets', 5)}+{last('BalanceSheets', 12)}"
    d["M55"] = "=" + "+".join(last("BalanceSheets", r) for r in (19, 20, 21, 24, 25))
    d["M57"] = f"={last('IncomeStatement', 23)}"
    # Assumption labels follow the projection years.
    for col in "FJN":
        d[f"{col}12"] = '="\'"&RIGHT($I$21,2)'
        d[f"{col}13"] = '="\'"&RIGHT($J$21,2)&"-\'"&RIGHT($M$21,2)'
        d[f"{col}14"] = '="\'"&RIGHT($I$21,2)'
        d[f"{col}15"] = '="\'"&RIGHT($M$21,2)'
    d["F12"] = None                                     # conservative revenue uses the base case
    # Starting assumptions from the company's own history instead of NVIDIA's numbers.
    d["K12"] = "=MIN(MEDIAN(E23:H23),H23)"              # base revenue growth, first projected year
    d["O12"] = "=K12*1.15"
    d["K15"] = "=AVERAGE(D29:H29)"                      # terminal EBIT margin
    d["G15"] = "=K15-0.03"
    d["O15"] = "=K15+0.02"
    d["I31"] = "=H29"                                   # margin starts at the last actual
    d["K18"] = "=WACC!D22"
    d["G18"] = "=K18+0.005"
    d["O18"] = "=K18-0.005"
    d["I40"] = "=AVERAGE(D40:H40)"                      # D&A % of sales
    d["I43"] = "=AVERAGE(D43:H43)"                      # CapEx % of sales
    d["I46"] = "=AVERAGE(D46:H46)"                      # ΔNWC % of sales

    w = wb["WACC"]
    w["D8"] = "=DCF!M55"
    w["D11"] = "=IFERROR(" + last("IncomeStatement", 15).replace("$H", "DCF!$H") + ",21%)"
    w["D13"] = "=GOOGLEFINANCE(DCF!C4)*DCF!M57"
    w["D17"] = '=IFERROR(GOOGLEFINANCE(DCF!C4,"beta"),1)'


# ---------------------------------------------------------------- export

def color(c):
    if c is None:
        return None
    if c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) == 8:
        return "#" + c.rgb[2:]
    return None     # theme colors: text = default black


BORDER = {"thin": "SOLID", "hair": "SOLID", "medium": "SOLID_MEDIUM", "thick": "SOLID_THICK",
          "dotted": "DOTTED", "dashed": "DASHED", "double": "DOUBLE", "mediumDashed": "DASHED"}


def style_of(cell):
    s = {}
    f = cell.font
    if f.b: s["b"] = True
    if f.i: s["i"] = True
    if f.u: s["u"] = True
    if f.name and f.name != "Calibri": s["font"] = f.name
    if f.sz and float(f.sz) != 11: s["sz"] = float(f.sz)
    fc = color(f.color)
    if fc and fc != "#000000": s["fc"] = fc
    if cell.fill is not None and cell.fill.fill_type == "solid":
        bg = color(cell.fill.fgColor)
        if bg and bg.upper() != "#FFFFFF": s["bg"] = bg
    if cell.number_format and cell.number_format != "General":
        s["nf"] = cell.number_format.replace("\\(", "(").replace("\\)", ")")
    if cell.alignment.horizontal in ("left", "center", "right"):
        s["ha"] = cell.alignment.horizontal.upper()
    b = cell.border
    for key, side in (("bt", b.top), ("bb", b.bottom), ("bl", b.left), ("br", b.right)):
        if side is not None and side.style in BORDER:
            s[key] = BORDER[side.style]
    return s


def export(wb):
    styles, style_index = [], {}
    sheets = []
    for ws in wb.worksheets:
        widths = {}
        for key, dim in ws.column_dimensions.items():
            if dim.width:
                for c in range(dim.min or column_index_from_string(key), (dim.max or column_index_from_string(key)) + 1):
                    if c <= MAX_COLS:
                        widths[get_column_letter(c)] = round(dim.width * DEFAULT_WIDTH_PX / EXCEL_DEFAULT_WIDTH)
        cells, max_r, max_c = [], 1, 1
        for row in ws.iter_rows(min_row=1, max_row=MAX_ROWS, max_col=MAX_COLS):
            for cell in row:
                v = unwrap_google(cell.value)
                st = style_of(cell)
                if v is None and not st:
                    continue
                item = {"a": cell.coordinate}
                if isinstance(v, str) and v.startswith("="):
                    item["f"] = v
                elif isinstance(v, bool):
                    item["v"] = "TRUE" if v else "FALSE"
                elif isinstance(v, (int, float)):
                    item["n"] = v
                elif v is not None and hasattr(v, "year"):
                    item["f"] = f"=DATE({v.year},{v.month},{v.day})"
                elif v is not None:
                    item["v"] = str(v)
                if st:
                    k = json.dumps(st, sort_keys=True)
                    if k not in style_index:
                        style_index[k] = len(styles)
                        styles.append(st)
                    item["s"] = style_index[k]
                cells.append(item)
                if v is not None:
                    max_r, max_c = max(max_r, cell.row), max(max_c, cell.column)
        sheets.append({
            "title": ws.title,
            "hideGridlines": ws.sheet_view.showGridLines is False,
            "rowCount": max(max_r + 10, 60),
            "columnCount": MAX_COLS,
            "columnWidths": widths,
            "merges": [str(r) for r in ws.merged_cells.ranges if r.max_row <= MAX_ROWS],
            "cells": cells,
        })
    return {"styles": styles, "sheets": sheets}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--universal", action="store_true")
    ap.add_argument("--json", required=True)
    ap.add_argument("--xlsx", dest="out_xlsx")
    ap.add_argument("--name", default="Valuation (universal)")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx)
    if a.universal:
        universalize(wb)
    data = {
        "name": a.name,
        "headerRow": HEADER_ROW,
        "yearColumns": {"first": YEAR_FIRST - 1, "count": YEAR_COUNT, "sheets": STATEMENTS},
        "placeholders": ["TICKER", "FY_FIRST", "TODAY", "FYEND"],
        **export(wb),
    }
    with open(a.json, "w") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    if a.out_xlsx:
        wb.save(a.out_xlsx)
    n = sum(len(s["cells"]) for s in data["sheets"])
    print(f"{a.json}: {len(data['sheets'])} sheets, {n} cells, {len(data['styles'])} styles")


if __name__ == "__main__":
    main()
